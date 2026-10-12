# -*- coding: utf-8 -*-
"""天地图物流运输距离与路径规划系统 - Streamlit 网页端
宏观审核精简版（含固定分类异常面板） + 微观单计划明细 + 本地地址缓存 + 智能Key降级休眠"""
import json
import os
import time
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
import xml.etree.ElementTree as ET
import cpca
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import requests
from datetime import datetime
import streamlit as st

# ==================== 页面配置 ====================
st.set_page_config(
    page_title="天地图物流运输距离与路径规划审核系统",
    page_icon="🚚",
    layout="wide"
)

# ==================== 全局配置与初始化 ====================
CACHE_FILE = "address_cache.json"
NO_PROXY = {"http": None, "https": None}  # 强制绕过本地代理直连物理网卡


# ==================== 本地地址缓存管理 ====================
def load_address_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            return {}
    return {}


def save_address_cache(cache_data):
    try:
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(cache_data, f, ensure_ascii=False, indent=4)
    except Exception as e:
        pass


# 全局内存缓存与线程锁
address_cache = load_address_cache()
cache_lock = threading.Lock()


# ==================== 智能 Key 管理器（支持降级与全局休眠） ====================
class SmartKeyManager:
    def __init__(self, keys):
        self.keys = keys
        self.cool_down = {k: 0 for k in keys}
        self.lock = threading.Lock()

    def get_valid_key(self):
        while True:
            with self.lock:
                now = time.time()
                for k in self.keys:
                    if now >= self.cool_down[k]:
                        return k
                min_wait = min(self.cool_down[k] - now for k in self.keys)

            sleep_time = max(min_wait, 1)
            time.sleep(sleep_time)

    def mark_key_limited(self, k):
        with self.lock:
            self.cool_down[k] = time.time() + 60  # 冷却 60 秒


# ==================== 错误分类映射函数（归纳为固定四大类） ====================
def classify_error(err_msg: str) -> str:
    if not err_msg or str(err_msg).strip() == "正常":
        return "正常"

    err_str = str(err_msg).lower()

    if any(kw in err_str for kw in ["cpca", "地址", "省市", "为空", "无效地址", "无法安全解析", "检验拦截"]):
        return "地址校验失败"
    elif any(kw in err_str for kw in ["限流", "freq", "limit", "冷却", "429", "频控"]):
        return "接口限流或频控"
    elif any(kw in err_str for kw in ["距离", "0", "无效", "路径规划失败", "明细获取失败", "未计算"]):
        return "路线规划或距离异常"
    else:
        return "其他网络或系统异常"


# ==================== 1. 地址校验与地理编码模块 ====================
def parse_and_validate_address(raw_address: str, role_name: str) -> dict:
    if not raw_address or not str(raw_address).strip():
        return {"success": False, "error": f"【{role_name}】地址为空或无效"}

    raw_str = str(raw_address).strip()
    df = cpca.transform([raw_str], pos_sensitive=True)
    row = df.iloc[0]

    province = str(row.get("省", "") or "")
    city = str(row.get("市", "") or "")
    district = str(row.get("区", "") or "")
    detail = str(row.get("地址", "") or "")

    if not province and not city:
        return {
            "success": False,
            "error": f"【{role_name}】地址原文 [{raw_str}] 未检测到省市关键字"
        }

    if province in ("上海市", "北京市", "天津市", "重庆市") and city in ("市辖区", province):
        city = ""

    parts = [province, city, district, detail]
    completed_address = "".join(p for p in parts if p)

    return {"success": True, "completed": completed_address}


def get_lat_lng(raw_address: str, role_name: str, key_manager, request_interval: int) -> dict:
    res = parse_and_validate_address(raw_address, role_name)
    if not res["success"]:
        return res

    completed_addr = res["completed"]

    with cache_lock:
        if completed_addr in address_cache:
            cached = address_cache[completed_addr]
            return {"success": True, "completed": completed_addr, "lon": cached["lon"], "lat": cached["lat"]}

    url = "http://api.tianditu.gov.cn/geocoder"

    for _ in range(3):
        tk = key_manager.get_valid_key()
        params = {"ds": json.dumps({"keyWord": completed_addr}, ensure_ascii=False), "tk": tk}

        try:
            time.sleep(request_interval)
            response = requests.get(url, params=params, proxies=NO_PROXY, timeout=20)
            data = response.json()

            if data.get("status") == "0":
                loc = data.get("location", {})
                lon, lat = loc.get("lon"), loc.get("lat")

                with cache_lock:
                    address_cache[completed_addr] = {"lon": lon, "lat": lat}
                    save_address_cache(address_cache)

                return {"success": True, "completed": completed_addr, "lon": lon, "lat": lat}

            msg = data.get('msg', '未知错误')
            if "限流" in msg or "freq" in msg.lower() or "limit" in msg.lower():
                key_manager.mark_key_limited(tk)
                continue
            else:
                return {"success": False, "error": f"天地图编码异常: {msg}"}
        except Exception as e:
            err_str = str(e)
            if "timed out" in err_str.lower():
                time.sleep(1)
                continue
            return {"success": False, "error": f"地理编码网络请求异常: {err_str}"}

    return {"success": False, "error": "地理编码失败：所有可用 Key 均尝试失败或限流"}


# ==================== 2. 路线规划与路段解析模块 ====================
def _fmt(p):
    return f"{p['lon']},{p['lat']}"


def _extract_toll_and_highway(root):
    statuses = [item.findtext("tollStatus") for item in root.findall("./routes/item")]
    statuses = [s for s in statuses if s is not None]
    toll = "收费" if any(s in ("1", "2") for s in statuses) else "免费" if statuses else "未知"

    names = [item.findtext("name") for item in root.findall(".//segment") or root.findall(".//item")]
    names = [n for n in names if n is not None]
    is_highway = "是" if any(any(hw in n for hw in ["高速", "绕城", "高架", "快速路"]) for n in names) else "否"

    return toll, is_highway


def get_drive_route(orig: dict, dest: dict, mids: list, key_manager, request_interval: int, max_retries=3) -> dict:
    url = "http://api.tianditu.gov.cn/drive"
    post = {"orig": _fmt(orig), "dest": _fmt(dest), "style": "0"}
    if mids:
        post["mid"] = ";".join([_fmt(m) for m in mids])

    for attempt in range(1, max_retries + 1):
        tk = key_manager.get_valid_key()
        params = {"postStr": json.dumps(post, ensure_ascii=False), "type": "search", "tk": tk}

        try:
            time.sleep(request_interval)
            response = requests.get(url, params=params, proxies=NO_PROXY, timeout=20)
            raw_content = response.content

            if raw_content.strip().startswith(b"<"):
                root = ET.fromstring(raw_content)
                distance = root.findtext("distance")
                duration = root.findtext("duration")

                if distance and float(distance) > 0:
                    toll, is_highway = _extract_toll_and_highway(root)
                    return {
                        "success": True,
                        "distance_km": float(distance),
                        "duration_min": round(float(duration) / 60, 1) if duration else 0.0,
                        "toll": toll,
                        "is_highway": is_highway
                    }
                else:
                    err_reason = f"天地图返回距离无效或为 0 (distance={distance})"
            else:
                text_snippet = raw_content[:100].decode('utf-8', errors='ignore')
                if "限流" in text_snippet or "freq" in text_snippet.lower():
                    key_manager.mark_key_limited(tk)
                    continue
                err_reason = f"非XML响应: {text_snippet!r}"
        except Exception as e:
            err_reason = f"网络请求异常: {str(e)}"

        if attempt < max_retries:
            time.sleep(1)
            continue
        else:
            return {"success": False, "error": f"整体路径规划失败（已重试 {max_retries} 次）: {err_reason}"}


def get_drive_segment(orig: dict, dest: dict, key_manager, request_interval: int, max_retries=3) -> dict:
    url = "http://api.tianditu.gov.cn/drive"
    post = {"orig": _fmt(orig), "dest": _fmt(dest), "style": "0"}

    for attempt in range(1, max_retries + 1):
        tk = key_manager.get_valid_key()
        params = {"postStr": json.dumps(post, ensure_ascii=False), "type": "search", "tk": tk}

        try:
            time.sleep(request_interval)
            response = requests.get(url, params=params, proxies=NO_PROXY, timeout=20)
            raw_content = response.content

            if raw_content.strip().startswith(b"<"):
                root = ET.fromstring(raw_content)
                distance = root.findtext("distance")
                duration = root.findtext("duration")

                if distance and float(distance) > 0:
                    toll, is_highway = _extract_toll_and_highway(root)
                    return {
                        "success": True,
                        "distance_km": float(distance),
                        "duration_min": round(float(duration) / 60, 1) if duration else 0.0,
                        "toll": toll,
                        "is_highway": is_highway
                    }
                else:
                    err_reason = f"单段路况距离无效或为 0 (distance={distance})"
            else:
                if "限流" in raw_content[:100].decode('utf-8', errors='ignore'):
                    key_manager.mark_key_limited(tk)
                    continue
                err_reason = "非XML响应"
        except Exception as e:
            err_reason = str(e)

        if attempt < max_retries:
            time.sleep(1)
            continue
        else:
            return {"success": False, "error": f"单段路况计算失败（已重试 {max_retries} 次）: {err_reason}"}


def drive_route_segmented(points, key_manager, request_interval):
    total_km, total_min = 0.0, 0.0
    for i in range(len(points) - 1):
        r = get_drive_segment(points[i], points[i + 1], key_manager, request_interval)
        if not r.get("success"):
            return r
        total_km += r["distance_km"]
        total_min += r["duration_min"]
    return {"success": True, "distance_km": total_km, "duration_min": total_min, "toll": "未知", "fallback": True}


# ==================== 3. 单条任务处理函数 ====================
def process_single_plan(idx, row, mid_cols, detail_output_dir, key_manager, request_interval):
    plan_id = row.get("计划编号", f"ROW_{idx + 1}")
    raw_orig = row.get("出发地")
    raw_dest = row.get("终点")
    row_data = {"计划编号": plan_id}

    try:
        nodes = []

        orig_res = get_lat_lng(raw_orig, "出发地", key_manager, request_interval)
        if not orig_res["success"]:
            raise ValueError(orig_res["error"])
        row_data["出发地"] = f"{raw_orig} ({orig_res['completed']})"
        nodes.append({
            "label": "出发地", "raw": raw_orig, "completed": orig_res["completed"],
            "lon": orig_res["lon"], "lat": orig_res["lat"]
        })

        mid_pts = []
        for mc in mid_cols:
            val = row.get(mc)
            if pd.notna(val) and str(val).strip():
                m_res = get_lat_lng(val, mc, key_manager, request_interval)
                if not m_res["success"]:
                    raise ValueError(m_res["error"])
                row_data[mc] = f"{val} ({m_res['completed']})"
                nodes.append({
                    "label": mc, "raw": val, "completed": m_res["completed"],
                    "lon": m_res["lon"], "lat": m_res["lat"]
                })
                mid_pts.append({"lon": m_res["lon"], "lat": m_res["lat"]})
            else:
                row_data[mc] = "-"

        dest_res = get_lat_lng(raw_dest, "终点", key_manager, request_interval)
        if not dest_res["success"]:
            raise ValueError(dest_res["error"])
        row_data["终点"] = f"{raw_dest} ({dest_res['completed']})"
        nodes.append({
            "label": "终点", "raw": raw_dest, "completed": dest_res["completed"],
            "lon": dest_res["lon"], "lat": dest_res["lat"]
        })

        orig_pt = {"lon": nodes[0]["lon"], "lat": nodes[0]["lat"]}
        dest_pt = {"lon": nodes[-1]["lon"], "lat": nodes[-1]["lat"]}

        route_res = get_drive_route(orig_pt, dest_pt, mid_pts, key_manager, request_interval)
        if not route_res.get("success"):
            all_pts = [orig_pt] + mid_pts + [dest_pt]
            route_res = drive_route_segmented(all_pts, key_manager, request_interval)

        if not route_res.get("success"):
            raise ValueError(route_res.get("error", "路径计算失败"))

        row_data["全程距离(公里)"] = route_res.get("distance_km")
        row_data["预计时长(分钟)"] = route_res.get("duration_min")
        row_data["异常状态"] = "正常"
        row_data["收费状态"] = route_res.get("toll")

        segment_records = []
        cum_dist = 0.0
        cum_time = 0.0

        for i in range(len(nodes) - 1):
            start_n = nodes[i]
            end_n = nodes[i + 1]
            seg = get_drive_segment(
                {"lon": start_n["lon"], "lat": start_n["lat"]},
                {"lon": end_n["lon"], "lat": end_n["lat"]},
                key_manager, request_interval
            )

            if not seg.get("success"):
                raise ValueError(f"微观路段明细获取失败: {seg.get('error')}")

            dist = seg.get("distance_km", 0.0)
            dur = seg.get("duration_min", 0.0)
            cum_dist += dist
            cum_time += dur
            avg_speed = round((dist / (dur / 60)), 1) if dur > 0 else 0.0

            segment_records.append({
                "路段序号": i + 1,
                "起点类型": start_n["label"],
                "起点原地址": start_n["raw"],
                "起点补全地址": start_n["completed"],
                "起点经纬度": f"{start_n['lon']},{start_n['lat']}",
                "终点类型": end_n["label"],
                "终点原地址": end_n["raw"],
                "终点补全地址": end_n["completed"],
                "终点经纬度": f"{end_n['lon']},{end_n['lat']}",
                "路段距离(公里)": dist,
                "路段预计时长(分钟)": dur,
                "路段平均时速(km/h)": avg_speed,
                "累计距离(公里)": round(cum_dist, 2),
                "累计时长(分钟)": round(cum_time, 1),
                "收费状态": seg.get("toll", "未知"),
                "是否高速/快速路": seg.get("is_highway", "否")
            })

        df_seg = pd.DataFrame(segment_records)
        detail_file_path = os.path.join(detail_output_dir, f"{plan_id}_路段明细.xlsx")
        df_seg.to_excel(detail_file_path, index=False, engine="openpyxl")

        wb_s = openpyxl.load_workbook(detail_file_path)
        ws_s = wb_s.active
        for col in ws_s.columns:
            max_l = max(sum(2 if ord(c) > 127 else 1 for c in str(cell.value or "")) for cell in col)
            ws_s.column_dimensions[get_column_letter(col[0].column)].width = max(max_l + 5, 15)
        wb_s.save(detail_file_path)

    except Exception as e:
        err_msg = str(e)
        row_data["出发地"] = f"{raw_orig} (检验拦截)" if pd.notna(raw_orig) else "-"
        for mc in mid_cols:
            val = row.get(mc)
            row_data[mc] = f"{val} (未计算)" if pd.notna(val) and str(val).strip() else "-"
        row_data["终点"] = f"{raw_dest} (检验拦截)" if pd.notna(raw_dest) else "-"
        row_data["全程距离(公里)"] = 0.0
        row_data["预计时长(分钟)"] = 0.0
        row_data["异常状态"] = err_msg
        row_data["收费状态"] = "-"

    return idx, row_data


# ==================== Streamlit 界面主逻辑 ====================
def main():
    st.title("🚚 天地图物流运输距离与路径规划审核系统")
    st.markdown("上传排车计划 Excel 文件，系统将自动执行地址清洗、智能缓存、双 Key 并发测算与四大异常分类。")

    with st.sidebar:
        st.header("⚙️ 系统配置")
        # 密钥输入框默认留空，必须由用户自行填写
        key1 = st.text_input("天地图 API Key 1", value="", type="password", placeholder="请输入您的天地图 Key")
        key2 = st.text_input("天地图 API Key 2", value="", type="password", placeholder="请输入备用 Key（可选）")

        request_interval = st.slider("请求间隔 (秒)", min_value=1, max_value=5, value=3)

    uploaded_file = st.file_uploader("上传排车计划 Excel 文件 (.xlsx)", type=["xlsx"])

    if uploaded_file is not None:
        try:
            df = pd.read_excel(uploaded_file)
            st.info(f"已成功加载排车计划文件，共计 **{len(df)}** 条记录。")

            with st.expander("查看原始数据预览"):
                st.dataframe(df.head(5))

            if st.button("🚀 开始自动化审核与路径规划", type="primary"):
                # 动态获取用户输入的有效 Key
                key_pool = [k.strip() for k in [key1, key2] if k and k.strip()]

                if not key_pool:
                    st.error("⚠️ 请先在左侧边栏输入至少一个有效的天地图 API Key 才能开始测算！")
                    st.stop()

                # 后台写死并发线程数为 2
                max_workers = 2
                key_manager = SmartKeyManager(key_pool)
                mid_cols = [c for c in df.columns if c.startswith("途径点")]

                result_dir = "05_结果"
                detail_output_dir = os.path.join(result_dir, "详细路段明细")
                os.makedirs(detail_output_dir, exist_ok=True)

                audit_records_dict = {}
                total_count = len(df)

                progress_bar = st.progress(0)
                status_text = st.empty()

                completed_count = 0

                with ThreadPoolExecutor(max_workers=max_workers) as executor:
                    futures = [
                        executor.submit(
                            process_single_plan, idx, row, mid_cols, detail_output_dir, key_manager, request_interval
                        )
                        for idx, row in df.iterrows()
                    ]

                    for future in as_completed(futures):
                        idx, row_data = future.result()
                        audit_records_dict[idx] = row_data
                        completed_count += 1

                        progress = completed_count / total_count
                        progress_bar.progress(progress)
                        status_text.text(
                            f"正在处理中... [{completed_count}/{total_count}] (计划编号: {row_data.get('计划编号')})")

                sorted_indices = sorted(audit_records_dict.keys())
                audit_records = [audit_records_dict[idx] for idx in sorted_indices]

                df_audit = pd.DataFrame(audit_records)
                df_audit["错误分类"] = df_audit["异常状态"].apply(classify_error)

                cols_order = ["计划编号", "出发地"] + mid_cols + ["终点", "全程距离(公里)", "预计时长(分钟)",
                                                                  "异常状态", "错误分类", "收费状态"]
                cols_order = [c for c in cols_order if c in df_audit.columns]
                df_audit = df_audit[cols_order]

                summary_out_path = os.path.join(result_dir, "排车计划审核清单_精简版.xlsx")

                with pd.ExcelWriter(summary_out_path, engine="openpyxl") as writer:
                    df_audit.to_excel(writer, sheet_name="全量汇总", index=False)
                    fixed_categories = ["地址校验失败", "接口限流或频控", "路线规划或距离异常", "其他网络或系统异常"]
                    for category in fixed_categories:
                        group_df = df_audit[df_audit["错误分类"] == category]
                        if len(group_df) > 0:
                            group_df.to_excel(writer, sheet_name=category, index=False)

                wb = openpyxl.load_workbook(summary_out_path)
                header_fill = PatternFill(start_color="365F91", end_color="365F91", fill_type="solid")
                header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                row_font = Font(name="Arial", size=10)
                zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
                white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
                border_thin = Border(
                    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
                )
                align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
                align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

                for ws in wb.worksheets:
                    ws.row_dimensions[1].height = 25
                    for col_num in range(1, ws.max_column + 1):
                        cell = ws.cell(row=1, column=col_num)
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = align_center
                        cell.border = border_thin

                    for row_num in range(2, ws.max_row + 1):
                        ws.row_dimensions[row_num].height = 24
                        is_zebra = (row_num % 2 == 0)
                        for col_num in range(1, ws.max_column + 1):
                            cell = ws.cell(row=row_num, column=col_num)
                            cell.font = row_font
                            cell.border = border_thin
                            cell.fill = zebra_fill if is_zebra else white_fill
                            header_name = ws.cell(row=1, column=col_num).value
                            if header_name in ["计划编号", "异常状态", "错误分类", "收费状态"]:
                                cell.alignment = align_center
                            else:
                                cell.alignment = align_left

                    for col in ws.columns:
                        max_len = 0
                        col_letter = get_column_letter(col[0].column)
                        for cell in col:
                            val_str = str(cell.value or "")
                            length = sum(2 if ord(c) > 127 else 1 for c in val_str)
                            if length > max_len:
                                max_len = length
                        ws.column_dimensions[col_letter].width = max(max_len + 6, 18)

                wb.save(summary_out_path)

                zip_path = os.path.join(result_dir, "详细路段明细打包.zip")
                with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
                    for root, dirs, files in os.walk(detail_output_dir):
                        for file in files:
                            full_path = os.path.join(root, file)
                            zipf.write(full_path, arcname=file)

                st.success("🎉 全部处理完成！审核结果已生成：")

                col1, col2 = st.columns(2)
                with col1:
                    with open(summary_out_path, "rb") as f:
                        st.download_button(
                            label="📥 下载宏观精简审核总表 (.xlsx)",
                            data=f,
                            file_name="排车计划审核清单_精简版.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        )
                with col2:
                    with open(zip_path, "rb") as f:
                        st.download_button(
                            label="📦 下载详细路段明细打包 (.zip)",
                            data=f,
                            file_name="详细路段明细打包.zip",
                            mime="application/zip"
                        )

        except Exception as e:
            st.error(f"文件处理过程中发生错误: {str(e)}")


if __name__ == "__main__":
    main()